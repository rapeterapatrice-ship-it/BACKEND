from rest_framework.decorators import api_view
from rest_framework.response import Response
from django.http import HttpResponse

from .models import Eleve, Classe, TypePaiement, MoisEcolage, CarnetFacturation, Presence
from .serializers import ClasseSerializer, TypePaiementSerializer, AnneeScolaireSerializer
from .helpers import get_annee_courante, MOIS_FULL
from .payment_logic import get_start_month, get_school_start_year, get_mois_facturables
from .auth import require_auth


def _get_eleves(annee_id, classe_id=None):
    qs = Eleve.objects.select_related('classe').filter(annee_id=annee_id)
    if classe_id:
        qs = qs.filter(classe_id=classe_id)
        qs = qs.order_by('nom', 'prenom')
    else:
        qs = qs.order_by('classe__niveau', 'classe__nom', 'nom', 'prenom')

    result = []
    for el in qs:
        mois_payes = list(MoisEcolage.objects.filter(
            eleve=el, annee_id=annee_id
        ).filter(type_paiement_id__in=[0]).values_list('mois_index', flat=True))
        result.append({
            'obj': el,
            'id': el.id, 'numero': el.numero, 'prenom': el.prenom, 'nom': el.nom,
            'classe_id': el.classe_id, 'classe_nom': el.classe.nom, 'niveau': el.classe.niveau,
            'tarif_ecolage': el.classe.tarif_ecolage,
            'fascicule_paye': el.fascicule_paye, 'location_payee': el.location_payee,
            'date_inscription': el.date_inscription,
            'moisPayes': mois_payes,
            'moisNoms': ', '.join(MOIS_FULL[i] for i in mois_payes) if mois_payes else '—',
        })
    return result


@api_view(['GET'])
@require_auth
def eleves_list_view(request):
    annee = get_annee_courante()
    classes = ClasseSerializer(Classe.objects.all().order_by('niveau', 'nom'), many=True).data
    niveaux = sorted(set(c['niveau'] for c in classes))
    types_paiement = TypePaiementSerializer(TypePaiement.objects.filter(actif=True).order_by('nom'), many=True).data

    if not annee:
        return Response({'eleves': [], 'classes': classes, 'classeId': None, 'annee': None,
                          'niveaux': niveaux, 'selectedNiveau': None, 'typesPaiement': []})

    classe_id = request.GET.get('classe_id') or None
    selected_niveau = request.GET.get('niveau') or None
    raw_eleves = _get_eleves(annee.id, classe_id)

    if selected_niveau:
        raw_eleves = [e for e in raw_eleves if e['niveau'] == selected_niveau]

    start_month = get_start_month()
    school_start_year = get_school_start_year(annee, start_month)

    eleves = []
    for el in raw_eleves:
        mois_facturables = get_mois_facturables(el['obj'], start_month, school_start_year)
        nb_mois_dus = len(mois_facturables)

        paiements_par_type = {}
        for tp in types_paiement:
            if tp['categorie'] == 'mensuel':
                mois_payes_type = list(MoisEcolage.objects.filter(
                    eleve_id=el['id'], annee_id=annee.id, type_paiement_id=tp['id']
                ).values_list('mois_index', flat=True))
                paiements_par_type[tp['id']] = {'moisPayes': mois_payes_type, 'nbMoisDus': nb_mois_dus, 'type': 'mensuel'}
            else:
                paye = CarnetFacturation.objects.filter(
                    eleve_id=el['id'], type_paiement_id=tp['id'], status='paye'
                ).filter(annee_id__in=[annee.id, None]).exists()
                paiements_par_type[tp['id']] = {'paye': paye, 'type': tp['categorie']}

        e = {k: v for k, v in el.items() if k != 'obj'}
        e['nbMoisDus'] = nb_mois_dus
        e['paiementsParType'] = paiements_par_type
        eleves.append(e)

    return Response({
        'eleves': eleves, 'classes': classes, 'classeId': int(classe_id) if classe_id else None,
        'annee': AnneeScolaireSerializer(annee).data, 'niveaux': niveaux,
        'selectedNiveau': selected_niveau, 'typesPaiement': types_paiement,
    })


@api_view(['GET'])
@require_auth
def eleves_export_view(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    annee = get_annee_courante()
    if not annee:
        return HttpResponse('Aucune année active', status=404)

    from .helpers import get_param
    classe_id = request.GET.get('classe_id') or None
    nom_etab = get_param('nom_etablissement', 'EduFinance')
    types_paiement = list(TypePaiement.objects.filter(actif=True).order_by('nom'))
    start_month = get_start_month()
    school_start_year = get_school_start_year(annee, start_month)

    raw_eleves = _get_eleves(annee.id, classe_id)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Élèves'

    base_headers = ['N°', 'Nom', 'Prénom', 'Classe', 'Niveau', 'Mois écolage payés', 'Nb mois écolage', 'Fascicule', 'Location']
    type_headers = [tp.nom for tp in types_paiement]
    all_headers = base_headers + type_headers
    total_cols = len(all_headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    c1 = ws.cell(row=1, column=1, value=f'{nom_etab} — Liste des élèves — A.S. {annee.libelle}')
    c1.font = Font(bold=True, size=13)
    c1.alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 22

    for i, h in enumerate(all_headers, start=1):
        cell = ws.cell(row=2, column=i, value=h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1A5C3A')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 18

    widths = [14, 20, 18, 12, 14, 50, 14, 12, 12] + [max(len(tp.nom) + 4, 14) for tp in types_paiement]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    thin = Side(style='thin', color='D1D5DB')
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for i, el in enumerate(raw_eleves):
        mois_facturables = get_mois_facturables(el['obj'], start_month, school_start_year)
        nb_mois_dus = len(mois_facturables)
        row_idx = 3 + i
        row_data = [el['numero'], el['nom'], el['prenom'], el['classe_nom'], el['niveau'],
                    el['moisNoms'], f"{len(el['moisPayes'])}/{nb_mois_dus}",
                    'Oui' if el['fascicule_paye'] else 'Non', 'Oui' if el['location_payee'] else 'Non']

        for tp in types_paiement:
            if tp.categorie == 'mensuel':
                mois_payes_type = MoisEcolage.objects.filter(
                    eleve_id=el['id'], annee_id=annee.id, type_paiement_id=tp.id
                ).count()
                row_data.append(f'{mois_payes_type}/{nb_mois_dus}')
            else:
                paye = CarnetFacturation.objects.filter(
                    eleve_id=el['id'], type_paiement_id=tp.id, status='paye'
                ).filter(annee_id__in=[annee.id, None]).exists()
                row_data.append('Oui' if paye else 'Non')

        for j, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=j, value=val)
            cell.border = border
            if i % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F4F8F6')

        for j in (8, 9):
            cell = ws.cell(row=row_idx, column=j)
            oui = cell.value == 'Oui'
            cell.font = Font(bold=True, color='166534' if oui else '6B7280')

        for ti in range(len(types_paiement)):
            cell = ws.cell(row=row_idx, column=10 + ti)
            val = str(cell.value or '')
            if val == 'Oui':
                cell.font = Font(bold=True, color='166534')
            elif val == 'Non':
                cell.font = Font(bold=True, color='6B7280')
            elif '/' in val:
                paid, total = (val.split('/') + ['0'])[:2]
                paid, total = int(paid), int(total or 0)
                color = '166534' if paid == total and total else ('B91C1C' if paid == 0 else '92400E')
                cell.font = Font(bold=True, color=color)

        mois_cell = ws.cell(row=row_idx, column=7)
        paid2, total2 = str(mois_cell.value or '0/0').split('/')
        paid2, total2 = int(paid2), int(total2)
        color2 = '166534' if paid2 == total2 else ('B91C1C' if paid2 == 0 else '92400E')
        mois_cell.font = Font(bold=True, color=color2)

    ws.freeze_panes = 'A3'

    fname = f"eleves_{annee.libelle.replace('-', '_')}" + ('_classe' if classe_id else '') + '.xlsx'
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{fname}"'
    wb.save(response)
    return response


@api_view(['DELETE', 'POST'])
@require_auth
def eleve_delete_view(request, eleve_id):
    from django.db import transaction
    from .models import LocationBook, LettrePdf, Entree
    try:
        with transaction.atomic():
            MoisEcolage.objects.filter(eleve_id=eleve_id).delete()
            LocationBook.objects.filter(eleve_id=eleve_id).delete()
            LettrePdf.objects.filter(eleve_id=eleve_id).delete()
            Entree.objects.filter(eleve_id=eleve_id).delete()
            CarnetFacturation.objects.filter(eleve_id=eleve_id).delete()
            Eleve.objects.filter(id=eleve_id).delete()
            Presence.objects.filter(eleve_id=eleve_id).delete()
    except Exception as e:
        return Response({'error': str(e)}, status=400)
    return Response({'ok': True})
